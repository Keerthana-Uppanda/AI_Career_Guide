# app.py

from typing import Set, Dict
import difflib
from io import BytesIO

import streamlit as st
import pandas as pd

# Optional imports
try:
    import pdfplumber
except Exception:
    pdfplumber = None

try:
    import docx
except Exception:
    docx = None

from fpdf import FPDF
from logic.career_matcher import load_careers, match_careers

# ===========================
# Page Config
# ===========================

st.set_page_config(
    page_title="AI Career Guide",
    page_icon="🎯",
    layout="wide"
)

# ===========================
# Helpers
# ===========================

def extract_text_from_pdf(file) -> str:
    text_parts = []
    with pdfplumber.open(file) as pdf:
        for page in pdf.pages:
            text_parts.append(page.extract_text() or "")
    return "\n".join(text_parts)

def extract_text_from_docx(file) -> str:
    document = docx.Document(file)
    return "\n".join(p.text for p in document.paragraphs)

def extract_resume_text(uploaded_file) -> str:
    name = uploaded_file.name.lower()
    if name.endswith(".pdf"):
        return extract_text_from_pdf(uploaded_file)
    elif name.endswith(".docx"):
        return extract_text_from_docx(uploaded_file)
    else:
        return uploaded_file.read().decode("utf-8", errors="ignore")

# ===========================
# Skill Logic
# ===========================

def get_skill_vocab(df: pd.DataFrame) -> Set[str]:
    skills = set()
    for cell in df["Key_Skills"].dropna():
        for s in str(cell).split(","):
            skills.add(s.strip().lower())
    return skills

def detect_resume_skills(text: str, skill_vocab: Set[str]) -> Set[str]:
    detected = set()
    text_lower = text.lower()
    for skill in skill_vocab:
        if skill in text_lower or difflib.get_close_matches(skill, text_lower.split(), n=1, cutoff=0.85):
            detected.add(skill)
    return detected

def score_career(skill_hits: Set[str], career_skill_cell: str) -> float:
    skills = {s.strip().lower() for s in str(career_skill_cell).split(",") if s.strip()}
    return len(skill_hits & skills) / len(skills) if skills else 0.0

# ===========================
# UI Components
# ===========================

def pill(text: str):
    st.markdown(
        f"""
        <span style="
            display:inline-block;
            padding:6px 14px;
            border-radius:999px;
            background:#4f46e5;
            color:white;
            font-size:0.85rem;
            margin:4px;">
            {text}
        </span>
        """,
        unsafe_allow_html=True
    )

def career_card(career: Dict, score: float):
    left, right = st.columns([4, 1])

    with left:
        st.markdown(
            f"""
            <div style="
                background:white;
                padding:20px;
                border-radius:18px;
                border:1px solid #e5e7eb;
                box-shadow:0 8px 20px rgba(0,0,0,0.08);
            ">
                <h3 style="margin-bottom:6px; color:#111827;">
                    💼 {career.get('Career','')}
                </h3>
                <p style="margin:0; color:#4b5563; font-size:15px;">
                    {career.get('Description','No description available')}
                </p>
            </div>
            """,
            unsafe_allow_html=True
        )

    with right:
        st.markdown(
            f"""
            <div style="
                background:#0f172a;
                color:white;
                padding:16px;
                border-radius:18px;
                text-align:center;
            ">
                <p style="margin:0; font-size:14px; color:#cbd5f5;">Match</p>
                <h2 style="margin:6px 0;">{score*100:.1f}%</h2>
            </div>
            """,
            unsafe_allow_html=True
        )
        st.progress(score)

    st.markdown("<br>", unsafe_allow_html=True)

# ===========================
# PDF Generation (FIXED)
# ===========================

def generate_pdf(user_name, detected_skills, filters, top_careers):
    pdf = FPDF()
    pdf.add_page()

    page_width = pdf.w - 2 * pdf.l_margin

    # ===== Title =====
    pdf.set_font("Arial", "B", 18)
    pdf.cell(0, 12, "AI Career Guide Report", ln=True, align="C")
    pdf.ln(4)

    pdf.set_font("Arial", size=11)
    pdf.cell(0, 8, f"Report for: {user_name}", ln=True, align="C")
    pdf.ln(8)

    # ===== Detected Skills =====
    pdf.set_font("Arial", "B", 14)
    pdf.cell(0, 10, "Detected Skills", ln=True)
    pdf.set_font("Arial", size=11)

    pdf.multi_cell(
        page_width,
        8,
        ", ".join(detected_skills) if detected_skills else "No skills detected"
    )
    pdf.ln(6)

    # ===== Filters =====
    pdf.set_font("Arial", "B", 14)
    pdf.cell(0, 10, "Filters Applied", ln=True)
    pdf.set_font("Arial", size=11)

    if filters:
        for k, v in filters.items():
            pdf.cell(0, 8, f"- {k.title()}: {v}", ln=True)
    else:
        pdf.cell(0, 8, "No filters applied", ln=True)
    pdf.ln(8)

    # ===== Career Recommendations =====
    pdf.set_font("Arial", "B", 14)
    pdf.cell(0, 10, "Career Recommendations", ln=True)
    pdf.ln(4)

    pdf.set_font("Arial", size=11)

    for idx, career in enumerate(top_careers, start=1):
        pdf.set_font("Arial", "B", 12)
        pdf.cell(
            0,
            8,
            f"{idx}. {career['Career']} (Match: {career['Score']*100:.1f}%)",
            ln=True
        )

        pdf.set_font("Arial", size=11)

        # Description
        desc = career.get("Description", "No description available")
        pdf.multi_cell(page_width, 7, f"Description: {desc}")

        # Key Skills
        skills = career.get("Key_Skills", "")
        if skills:
            pdf.multi_cell(page_width, 7, f"Key Skills: {skills}")

        # Subjects
        subjects = career.get("Subjects", "")
        if subjects:
            pdf.multi_cell(page_width, 7, f"Subjects: {subjects}")

        pdf.ln(5)

    # ===== Suggested Skills to Learn =====
    pdf.set_font("Arial", "B", 14)
    pdf.cell(0, 10, "Suggested Skills to Learn", ln=True)
    pdf.set_font("Arial", size=11)

    suggested = set()
    for career in top_careers:
        req = {
            s.strip().lower()
            for s in str(career.get("Key_Skills", "")).split(",")
            if s.strip()
        }
        suggested |= (req - set(detected_skills))

    if suggested:
        pdf.multi_cell(page_width, 8, ", ".join(sorted(suggested)))
    else:
        pdf.cell(0, 8, "You already match most required skills!", ln=True)

    # ===== Footer =====
    pdf.ln(10)
    pdf.set_font("Arial", size=9)
    pdf.cell(0, 8, "Generated by AI Career Guide", align="C")

    return pdf


# ===========================
# HERO SECTION
# ===========================

st.markdown(
    """
    <div style="text-align:center; padding:20px 0;">
        <h1>🎯 AI Career Guide</h1>
        <p style="font-size:18px; color:#6b7280;">
            Upload your resume and discover careers that truly fit you
        </p>
    </div>
    """,
    unsafe_allow_html=True
)

# ===========================
# SIDEBAR
# ===========================

with st.sidebar:
    st.header("⚙️ Filters")
    education = st.selectbox("Education Level", ["", "Diploma", "Bachelor's", "Master's", "PhD"])
    top_n = st.slider("Number of Careers", 1, 10, 5)

# ===========================
# MAIN LOGIC
# ===========================

df = load_careers("data/careers.csv")
uploaded = st.file_uploader("📄 Upload your Resume", type=["pdf", "docx", "txt"])

if uploaded:
    resume_text = extract_resume_text(uploaded)
    skill_vocab = get_skill_vocab(df)
    hits = detect_resume_skills(resume_text, skill_vocab)

    st.divider()
    st.subheader("🔍 Detected Skills")
    for s in sorted(hits):
        pill(s)

    user_input = {}
    if education:
        user_input["education"] = education

    filtered = match_careers(user_input, df)

    scores = []
    for idx, row in filtered.iterrows():
        scores.append((idx, score_career(hits, row["Key_Skills"])) )

    scores.sort(key=lambda x: x[1], reverse=True)
    top_rows = scores[:top_n]

    st.divider()
    st.subheader("🧭 Recommended Careers")

    top_career_list = []
    for i, s in top_rows:
        row = filtered.loc[i].to_dict()
        row["Score"] = s
        career_card(row, s)
        top_career_list.append(row)

    st.divider()
    st.subheader("📄 Download Report")

    pdf = generate_pdf("Resume Analysis", sorted(hits), user_input, top_career_list)
    buffer = BytesIO()
    pdf.output(buffer)

    st.download_button(
        "⬇️ Download PDF Report",
        data=buffer.getvalue(),
        file_name="career_report.pdf",
        mime="application/pdf"
    )

else:
    st.info("⬆️ Upload your resume to get started")
